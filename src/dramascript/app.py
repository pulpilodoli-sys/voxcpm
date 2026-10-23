import json
import os
import sys
import threading
from pathlib import Path

from PySide6.QtCore import QThread, Signal, QUrl, QTimer
from PySide6.QtGui import QDesktopServices, QTextDocument, QFont
from PySide6.QtPrintSupport import QPrinter
from PySide6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout,
    QHBoxLayout, QFormLayout, QLineEdit, QPushButton, QFileDialog, QComboBox, QSpinBox,
    QPlainTextEdit, QTabWidget, QProgressBar, QLabel, QMessageBox)
from docx import Document
from .engine import process, Cancelled
from .review import review


class Worker(QThread):
    update = Signal(int, str)
    done = Signal(str, object)
    error = Signal(str)

    def __init__(self, video, project, config, reviewing=False):
        super().__init__()
        self.args = video, project, config
        self.reviewing = reviewing
        self.stop = threading.Event()

    def run(self):
        try:
            if self.reviewing:
                text, data = review(self.args[1], self.args[2], self.update.emit, self.stop.is_set)
            else:
                text, data = process(*self.args, self.update.emit, self.stop.is_set)
            self.done.emit(text, data)
        except Cancelled as e:
            self.error.emit(str(e))
        except Exception as e:
            # Avoid reflecting any user-supplied key through exception messages.
            message = str(e)
            key = self.args[2].get('key', '')
            self.error.emit(message.replace(key, '[密钥已隐藏]') if key else message)


class Window(QMainWindow):
    def __init__(self):
        super().__init__()
        self.worker = None
        self.setWindowTitle('叙幕 · 中文剧情剧本工作台')
        self.resize(1150, 850)
        root = QWidget()
        self.setCentralWidget(root)
        layout = QVBoxLayout(root)
        title = QLabel('叙幕  /  从视频到中文剧情剧本')
        title.setStyleSheet('font-size:24px; font-weight:bold; padding:12px 0')
        layout.addWidget(title)
        layout.addWidget(QLabel('剧情剧本 · 人物连续性 · 无对白动作描写 · Windows / macOS'))
        form = QFormLayout()
        layout.addLayout(form)
        self.video = QLineEdit()
        self.project = QLineEdit()
        for label, field, slot in [('视频文件', self.video, self.pick_video), ('项目文件夹', self.project, self.pick_project)]:
            row = QHBoxLayout()
            row.addWidget(field)
            button = QPushButton('选择…')
            button.clicked.connect(slot)
            row.addWidget(button)
            form.addRow(label, row)
        self.mode = QComboBox()
        self.mode.addItems(['云端', '本地'])
        form.addRow('分析方式', self.mode)
        self.url = QLineEdit('https://api.openai.com/v1')
        self.model = QLineEdit('gpt-4.1')
        self.audio = QLineEdit('whisper-1')
        self.key = QLineEdit(os.environ.get('DRAMASCRIPT_API_KEY', ''))
        self.key.setEchoMode(QLineEdit.Password)
        self.key.setPlaceholderText('仅保存在内存；不会写入项目')
        form.addRow('兼容 OpenAI 的接口地址', self.url)
        form.addRow('支持图像输入的模型', self.model)
        form.addRow('转写模型', self.audio)
        form.addRow('云端 API 密钥', self.key)
        self.mode.currentTextChanged.connect(self.change_mode)
        self.notice = QLabel('云端模式会向所填接口上传音频与抽样画面，可能产生费用。点击开始即使用该服务。')
        self.notice.setWordWrap(True)
        layout.addWidget(self.notice)
        actions = QHBoxLayout()
        self.start = QPushButton('开始 / 继续分析')
        self.start.clicked.connect(self.analyze)
        self.review_button = QPushButton('全片剧情复核 / 人物档案')
        self.review_button.clicked.connect(lambda: self.analyze(reviewing=True))
        self.pause = QPushButton('暂停（当前请求结束后）')
        self.pause.setEnabled(False)
        self.pause.clicked.connect(lambda: self.worker.stop.set())
        self.open = QPushButton('打开已有项目')
        self.open.clicked.connect(self.open_project)
        for b in (self.start, self.pause, self.open, self.review_button):
            actions.addWidget(b)
        layout.addLayout(actions)
        self.progress = QProgressBar()
        layout.addWidget(self.progress)
        self.status = QLabel('就绪。自动生成内容需对照视频核实；抽样分析可能遗漏短暂动作。')
        layout.addWidget(self.status)
        self.tabs = QTabWidget()
        self.script = QPlainTextEdit()
        self.script.setPlaceholderText('中文剧情剧本将显示在此。可编辑，并保存为项目中的独立修订稿。')
        self.evidence = QPlainTextEdit()
        self.evidence.setReadOnly(True)
        self.tabs.addTab(self.script, '中文剧本 · 可编辑')
        self.tabs.addTab(self.evidence, '剧情证据 / 人物 / 待核实')
        self.report = QPlainTextEdit()
        self.report.setReadOnly(True)
        self.tabs.addTab(self.report, '全片复核 / 人物档案')
        self.segments = QComboBox()
        self.segments.currentIndexChanged.connect(self.select_segment)
        evidence_row = QHBoxLayout()
        evidence_row.addWidget(QLabel('核对片段'))
        evidence_row.addWidget(self.segments, 1)
        self.seek_seconds = QSpinBox()
        self.seek_seconds.setRange(0, 604800)
        self.seek_seconds.setSuffix(' 秒')
        evidence_row.addWidget(self.seek_seconds)
        self.preview_button = QPushButton('提取该时间画面')
        self.preview_button.clicked.connect(self.preview_frame)
        evidence_row.addWidget(self.preview_button)
        layout.addLayout(evidence_row)
        layout.addWidget(self.tabs, 1)
        bottom = QHBoxLayout()
        for label, slot in [('保存修订稿', self.save_edit), ('导出 DOCX', self.export_docx), ('导出 PDF', self.export_pdf), ('查看原视频', self.open_video)]:
            b = QPushButton(label)
            b.clicked.connect(slot)
            bottom.addWidget(b)
        layout.addLayout(bottom)
        self.setStyleSheet('QWidget {font-size:14px;} QMainWindow {background:#f3f5f8;} QPushButton {padding:8px;} QPlainTextEdit {background:white; padding:10px;}')

    def change_mode(self, mode):
        local = mode == '本地'
        self.url.setText('http://localhost:11434/v1' if local else 'https://api.openai.com/v1')
        self.model.setText('qwen2.5vl:7b' if local else 'gpt-4.1')
        self.audio.setText('small' if local else 'whisper-1')
        self.key.setEnabled(not local)
        self.notice.setText('本地：需自行启动 Ollama 视觉模型并安装本地转写依赖。首次下载模型需要网络；推理使用本机 CPU，长视频可能耗时较久。' if local else '云端模式会向所填接口上传音频与抽样画面，可能产生费用。点击开始即使用该服务。')

    def pick_video(self):
        if self.worker and self.worker.isRunning():
            return
        path, _ = QFileDialog.getOpenFileName(self, '选择长视频', '', '视频 (*.mp4 *.mkv *.mov *.avi *.webm);;所有文件 (*)')
        if path:
            self.video.setText(path)

    def pick_project(self):
        if self.worker and self.worker.isRunning():
            return
        path = QFileDialog.getExistingDirectory(self, '选择项目文件夹（可为空文件夹）')
        if path:
            self.project.setText(path)

    def analyze(self, checked=False, reviewing=False):
        if (not reviewing and not Path(self.video.text()).is_file()) or not self.project.text():
            QMessageBox.warning(self, '需要输入', '请选择视频文件和项目文件夹。')
            return
        if self.script.document().isModified():
            QMessageBox.warning(self, '请先保存', '请先保存修订稿，分析完成会载入新的草稿。')
            return
        config = dict(mode=self.mode.currentText(), url=self.url.text().strip(), model=self.model.text().strip(),
                      audio_model=self.audio.text().strip(), key=self.key.text() if self.mode.currentText() == '云端' else '')
        if not config['url'].startswith(('https://', 'http://localhost:', 'http://127.0.0.1:')):
            QMessageBox.warning(self, '接口地址', '远程接口必须使用 HTTPS；本机服务允许 HTTP。')
            return
        if reviewing and not (Path(self.project.text()) / 'project.json').exists():
            QMessageBox.warning(self, '请先分析', '请先打开已有项目或完成视频分析。')
            return
        self.worker = Worker(self.video.text(), self.project.text(), config, reviewing)
        self.worker.update.connect(self.update_progress)
        self.worker.done.connect(self.complete)
        self.worker.error.connect(self.status.setText)
        self.worker.finished.connect(self.finished)
        self.start.setEnabled(False)
        self.review_button.setEnabled(False)
        self.script.setReadOnly(True)
        self.open.setEnabled(False)
        self.pause.setEnabled(True)
        for field in (self.video, self.project, self.mode, self.url, self.model, self.audio, self.key):
            field.setEnabled(False)
        self.worker.start()

    def update_progress(self, value, text):
        self.progress.setValue(value)
        self.status.setText(text)

    def finished(self):
        for field in (self.video, self.project, self.mode, self.url, self.model, self.audio, self.key):
            field.setEnabled(True)
        self.key.setEnabled(self.mode.currentText() == "云端")
        self.start.setEnabled(True)
        self.review_button.setEnabled(True)
        self.script.setReadOnly(False)
        self.open.setEnabled(True)
        self.pause.setEnabled(False)

    def complete(self, text, data):
        if self.worker and self.worker.reviewing:
            self.report.setPlainText(text)
            self.tabs.setCurrentWidget(self.report)
            self.status.setText('证据报告已保存；人物线索只通过格式与引用校验，身份和性格仍需确认。' if data.get('review_status', {}).get('schema_and_references_validated') else '证据报告已保存；部分模型线索未通过校验，尚不能确认人物档案。')
            return
        self.script.setPlainText(text)
        self.script.document().setModified(False)
        self.show_evidence(data)
        self.progress.setValue(100)
        self.status.setText('草稿已生成并保存。请核实人物、对白与跨场景关系后导出。')

    def show_evidence(self, data):
        self.segment_data = data['chunks']
        self.segments.clear()
        self.segments.addItems([f'{c["start"]:.1f}–{c["end"]:.1f} 秒' for c in data['chunks']])
        report = Path(self.project.text()) / '全片审校报告.txt'
        self.report.setPlainText(report.read_text(encoding='utf-8') if report.exists() else '尚未生成全片审校报告。')
        self.evidence.setPlainText('\n\n'.join(f'【{c["start"]:.1f}–{c["end"]:.1f} 秒】\n{c["analysis"]}\n\n原始转写：\n{c["transcript"]}' for c in data['chunks']))

    def open_project(self):
        if self.script.document().isModified():
            QMessageBox.warning(self, '请先保存', '请先保存修订稿，再打开其他项目。')
            return
        path = QFileDialog.getExistingDirectory(self, '打开项目')
        if not path:
            return
        try:
            folder = Path(path)
            data = json.loads((folder / 'project.json').read_text(encoding='utf-8'))
            self.project.setText(path)
            self.video.setText(data['video'])
            self.mode.setCurrentText(data['config']['mode'])
            self.url.setText(data['config']['url'])
            self.model.setText(data['config']['model'])
            self.audio.setText(data['config']['audio_model'])
            revision = folder / '剧本修订稿.txt'
            self.script.setPlainText(revision.read_text(encoding='utf-8') if revision.exists() else '\n\n'.join(c['script'] for c in data['chunks']))
            self.script.document().setModified(False)
            self.show_evidence(data)
            self.status.setText('项目已打开；可编辑或继续未完成的分析。')
        except Exception as e:
            QMessageBox.warning(self, '打开失败', str(e))

    def save_edit(self):
        if not self.project.text():
            return
        try:
            folder = Path(self.project.text())
            folder.mkdir(parents=True, exist_ok=True)
            (folder / '剧本修订稿.txt').write_text(self.script.toPlainText(), encoding='utf-8')
            self.script.document().setModified(False)
            self.status.setText('修订稿已保存；自动分析草稿单独保留。')
        except OSError as e:
            QMessageBox.warning(self, '保存失败', str(e))

    def export_docx(self):
        path, _ = QFileDialog.getSaveFileName(self, '导出 DOCX', '中文剧情剧本.docx', 'Word 文档 (*.docx)')
        if path:
            try:
                doc = Document()
                doc.add_heading('中文剧情剧本', 0)
                for line in self.script.toPlainText().splitlines():
                    doc.add_paragraph(line)
                doc.save(path)
                self.status.setText('DOCX 已导出。')
            except Exception as e:
                QMessageBox.warning(self, '导出失败', str(e))

    def export_pdf(self):
        path, _ = QFileDialog.getSaveFileName(self, '导出 PDF', '中文剧情剧本.pdf', 'PDF 文档 (*.pdf)')
        if path:
            try:
                printer = QPrinter(QPrinter.HighResolution)
                printer.setOutputFormat(QPrinter.PdfFormat)
                printer.setOutputFileName(path)
                document = QTextDocument()
                document.setDefaultFont(QFont('Noto Sans CJK SC', 11))
                document.setPlainText(self.script.toPlainText())
                document.print_(printer)
                self.status.setText('PDF 已导出；请确认本机具备中文字体。')
            except Exception as e:
                QMessageBox.warning(self, '导出失败', str(e))

    def select_segment(self, index):
        if index >= 0 and hasattr(self, 'segment_data'):
            self.seek_seconds.setValue(int(self.segment_data[index]['start']))
            self.evidence.moveCursor(self.evidence.textCursor().MoveOperation.Start)
            self.evidence.find(f'【{self.segment_data[index]["start"]:.1f}–')

    def preview_frame(self):
        from .engine import command
        import tempfile
        if not Path(self.video.text()).is_file():
            QMessageBox.warning(self, '视频不可用', '请选择原视频。')
            return
        try:
            folder = Path(tempfile.mkdtemp(prefix='dramascript-preview-'))
            frame = folder / 'preview.jpg'
            command(['ffmpeg', '-v', 'error', '-ss', str(self.seek_seconds.value()),
                     '-i', self.video.text(), '-frames:v', '1', '-vf', 'scale=1280:-2,format=yuvj420p', str(frame)])
            if not frame.exists():
                raise RuntimeError('该时间无画面，请选择视频范围内的时间。')
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(frame)))
            self.status.setText(f'已打开 {self.seek_seconds.value()} 秒画面；完整动作请对照原视频。')
        except Exception as e:
            QMessageBox.warning(self, '提取失败', str(e))

    def open_video(self):
        if Path(self.video.text()).is_file():
            QDesktopServices.openUrl(QUrl.fromLocalFile(self.video.text()))

    def closeEvent(self, event):
        if self.worker and self.worker.isRunning():
            self.worker.stop.set()
            QMessageBox.information(self, '正在暂停', '请等待当前模型请求结束后关闭窗口。')
            event.ignore()
        elif self.script.document().isModified():
            QMessageBox.information(self, '未保存的修改', '请先保存修订稿，再关闭窗口。')
            event.ignore()
        else:
            event.accept()


def main():
    if '--check-media' in sys.argv:
        from .runtime import check_media
        check_media()
        return
    app = QApplication(sys.argv)
    window = Window()
    window.show()
    if '--smoke-test' in sys.argv:
        QTimer.singleShot(500, app.quit)
    result = app.exec()
    import shiboken6
    if shiboken6.isValid(window):
        shiboken6.delete(window)
    sys.exit(result)


if __name__ == '__main__':
    main()
